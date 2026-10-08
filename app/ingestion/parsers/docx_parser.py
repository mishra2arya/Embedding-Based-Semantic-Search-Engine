"""DOCX document parser using python-docx."""

from __future__ import annotations

import io

from docx import Document

from app.core.exceptions import DocumentParsingError
from app.ingestion.cleaners.normalizer import TextNormalizer
from app.ingestion.loaders.file_loader import RawDocument
from app.ingestion.parsers.base import BaseParser, ParsedDocument


class DOCXParser(BaseParser):
    """Parses Microsoft Word (.docx) documents."""

    def parse(self, raw_doc: RawDocument) -> ParsedDocument:
        try:
            doc = Document(io.BytesIO(raw_doc.content))
        except Exception as e:
            raise DocumentParsingError(f"Corrupt or unreadable DOCX {raw_doc.filename}: {e}") from e

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        normalized_text = TextNormalizer.normalize(full_text)

        title = raw_doc.filename
        if paragraphs:
            title = paragraphs[0][:120]

        metadata = dict(raw_doc.metadata)
        metadata.update(
            {
                "parser": "DOCXParser",
                "format": "docx",
                "paragraph_count": len(paragraphs),
            }
        )

        return ParsedDocument(
            title=title,
            text=normalized_text,
            filename=raw_doc.filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            source=raw_doc.source,
            tenant_id=raw_doc.tenant_id,
            metadata=metadata,
        )
