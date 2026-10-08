"""Unit tests for document loaders and specialized parsers."""

from __future__ import annotations

import io
import json

import docx
import pypdf
import pytest

from app.core.exceptions import DocumentParsingError, PayloadTooLargeError
from app.ingestion.loaders.file_loader import FileSystemLoader, RawDocument
from app.ingestion.parsers.docx_parser import DOCXParser
from app.ingestion.parsers.json_parser import JSONParser
from app.ingestion.parsers.pdf_parser import PDFParser


def test_file_system_loader(tmp_path):
    loader = FileSystemLoader(max_size_mb=1, allowed_extensions=[".txt", ".json"])

    # Non-existent path
    with pytest.raises(FileNotFoundError):
        list(loader.load(tmp_path / "nonexistent"))

    # Create files
    file_txt = tmp_path / "doc.txt"
    file_txt.write_text("Text content", encoding="utf-8")

    file_json = tmp_path / "doc.json"
    file_json.write_text('{"key": "value"}', encoding="utf-8")

    file_ignored = tmp_path / "ignore.xyz"
    file_ignored.write_text("Ignored", encoding="utf-8")

    # Load single file
    single_docs = list(loader.load(file_txt))
    assert len(single_docs) == 1
    assert single_docs[0].content == b"Text content"
    assert single_docs[0].filename == "doc.txt"

    # Load directory
    dir_docs = list(loader.load(tmp_path))
    assert len(dir_docs) == 2
    filenames = {d.filename for d in dir_docs}
    assert filenames == {"doc.txt", "doc.json"}

    # File exceeding size limit
    big_file = tmp_path / "big.txt"
    big_file.write_bytes(b"x" * (2 * 1024 * 1024))
    with pytest.raises(PayloadTooLargeError):
        list(loader.load(big_file))


def test_json_parser_variants():
    parser = JSONParser()

    # 1. JSONL parsing
    jsonl_bytes = b'{"text": "Line 1"}\n\n{"content": "Line 2"}\n{"invalid": \n{"body": "Line 3"}'
    raw_jsonl = RawDocument(content=jsonl_bytes, filename="stream.jsonl", mime_type="application/x-ndjson")
    parsed_jsonl = parser.parse(raw_jsonl)
    assert "Line 1" in parsed_jsonl.text
    assert "Line 2" in parsed_jsonl.text
    assert "Line 3" in parsed_jsonl.text

    # 2. JSON dict without standard text field (key: value pairs)
    dict_bytes = json.dumps({"service": "vector-search", "version": "1.0"}).encode("utf-8")
    raw_dict = RawDocument(content=dict_bytes, filename="config.json")
    parsed_dict = parser.parse(raw_dict)
    assert "service: vector-search" in parsed_dict.text

    # 3. JSON with rich metadata
    meta_doc = {
        "title": "Rich Doc",
        "text": "Body content here",
        "author": "Alice",
        "category": "ai",
        "language": "en",
        "tags": ["ml", "rag"],
        "created_at": "2026-01-01",
        "metadata": {"custom_score": 99},
    }
    raw_meta = RawDocument(content=json.dumps(meta_doc).encode("utf-8"), filename="meta.json")
    parsed_meta = parser.parse(raw_meta)
    assert parsed_meta.title == "Rich Doc"
    assert parsed_meta.metadata["author"] == "Alice"
    assert parsed_meta.metadata["category"] == "ai"
    assert parsed_meta.metadata["tags"] == ["ml", "rag"]
    assert parsed_meta.metadata["custom_score"] == 99

    # 4. JSON array of items
    array_bytes = json.dumps([{"text": "Item A"}, {"content": "Item B"}, "Primitive item"]).encode("utf-8")
    raw_arr = RawDocument(content=array_bytes, filename="arr.json")
    parsed_arr = parser.parse(raw_arr)
    assert "Item A" in parsed_arr.text
    assert "Item B" in parsed_arr.text
    assert "Primitive item" in parsed_arr.text

    # 5. Malformed JSON fallback
    raw_bad = RawDocument(content=b"not json { [", filename="bad.json")
    parsed_bad = parser.parse(raw_bad)
    assert "not json" in parsed_bad.text


def test_docx_parser():
    parser = DOCXParser()

    # Valid docx document
    doc = docx.Document()
    doc.add_paragraph("First Header Paragraph")
    doc.add_paragraph("Second body paragraph of the document.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    raw_doc = RawDocument(
        content=docx_bytes,
        filename="sample.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    parsed = parser.parse(raw_doc)
    assert parsed.title == "First Header Paragraph"
    assert "Second body paragraph" in parsed.text
    assert parsed.metadata["paragraph_count"] == 2
    assert parsed.metadata["format"] == "docx"

    # Corrupted docx
    raw_corrupt = RawDocument(content=b"corrupt non-docx bytes", filename="bad.docx")
    with pytest.raises(DocumentParsingError):
        parser.parse(raw_corrupt)


def test_pdf_parser():
    parser = PDFParser()

    # Create a small valid PDF in-memory using pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_metadata({"/Title": "PDF Test Title", "/Author": "Test Author"})
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    raw_pdf = RawDocument(content=pdf_bytes, filename="doc.pdf", mime_type="application/pdf")
    parsed = parser.parse(raw_pdf)
    assert parsed.title == "PDF Test Title"
    assert parsed.metadata["author"] == "Test Author"
    assert parsed.metadata["page_count"] == 1
    assert parsed.metadata["format"] == "pdf"

    # Corrupt PDF
    raw_corrupt = RawDocument(content=b"not a real pdf content", filename="corrupt.pdf")
    with pytest.raises(DocumentParsingError):
        parser.parse(raw_corrupt)
